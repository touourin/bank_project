"""Read back exports and independently reconcile their meaning to preserved sources."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from decimal import Decimal

import openpyxl

from .split_sources import BASES, digest, write_json
from .summarize import OUT
from .summary_common import Sources


def close(actual, expected):
    assert actual is not None and abs(Decimal(str(actual)) - Decimal(str(expected))) <= Decimal(
        "0.00000001"
    ), (actual, expected)


def verify(version, label):
    folder = OUT / version / label
    spec = json.loads((folder / "input.json").read_text())
    assert digest(folder / "before.xlsx") == spec["source_sha256"]
    assert digest(BASES[version] / f"{label}.xlsx") == spec["source_sha256"], (
        "Input changed during processing"
    )
    workbook = openpyxl.load_workbook(folder / "after.xlsx", read_only=False, data_only=False)
    assert workbook.sheetnames == ["数据", "字段说明"]
    fields = spec["fields"]
    keys = [f["key"] for f in fields]
    ws = workbook["数据"]
    actual = list(ws.values)
    assert list(actual[0]) == keys
    assert len(actual) == len(spec["rows"]) + 1
    assert len(ws.tables) == 1 and ws.freeze_panes == "C2"
    for row_no, (found, expected) in enumerate(zip(actual[1:], spec["rows"], strict=True), start=2):
        for col_no, (value, want, field) in enumerate(
            zip(found, expected, fields, strict=True), start=1
        ):
            cell = ws.cell(row_no, col_no)
            assert cell.data_type != "f", "Import fixture must be static"
            if want is None:
                assert value is None, (row_no, col_no, value)
            elif field["kind"] == "text":
                assert isinstance(value, str) and value == want, (row_no, col_no, value, want)
                assert cell.number_format == "@"
            else:
                assert isinstance(value, (float, int))
                close(value, want)
    notes = list(workbook["字段说明"].iter_rows(min_row=10, values_only=True))
    assert len(notes) == len(keys) and [r[0] for r in notes] == keys
    assert all(r[1] and r[4] and r[5] for r in notes)
    rows = [dict(zip(keys, r, strict=True)) for r in actual[1:]]
    source = Sources(folder / "before.xlsx")
    result = dict(
        rows=len(rows),
        fields=len(keys),
        checked_cells=len(rows) * len(keys),
        linked_customers=sum(r["cust_ind"] is not None for r in rows),
        source_tables=len(source.tables),
        source_records=sum(map(len, source.tables.values())),
        source_sha256=spec["source_sha256"],
        output_sha256=digest(folder / "after.xlsx"),
    )
    if label == "征信":
        assert len({(r["subject_type"], r["report_id"]) for r in rows}) == len(rows)
        assert result["linked_customers"] == 1000
        categories = Counter(r["subject_type"] for r in rows)
        result["subject_types"] = dict(categories)
        assert categories == {"企业及其他组织": 1010, "个人": 100}
        groups = {False: set(), True: set()}
        for r in rows:
            groups[r["subject_type"] == "个人"].add(r["report_id"])
        # Every detail record belongs to a represented report of the right subject type.
        for table, records in source.tables.items():
            personal = table.startswith("PBCPC")
            key = "PA01AI01" if personal else "EA01AI01"
            assert all(r[key] in groups[personal] for r in records), table
        conflicts, unknown, multi = 0, 0, 0
        financial_pairs = {
            "assets": (
                "PBCEC_EG02AB_DEBTINFO2007",
                "EG02BJ31",
                "PBCEC_EG01AB_DEBTINFO2002",
                "EG01BJ44",
            ),
            "liabilities": (
                "PBCEC_EG02AB_DEBTINFO2007",
                "EG02BJ53",
                "PBCEC_EG01AB_DEBTINFO2002",
                "EG01BJ68",
            ),
            "equity": (
                "PBCEC_EG02AB_DEBTINFO2007",
                "EG02BJ59",
                "PBCEC_EG01AB_DEBTINFO2002",
                "EG01BJ86",
            ),
            "net_profit": (
                "PBCEC_EG04AB_PROFITINFO2007",
                "EG04BJ17",
                "PBCEC_EG03AB_PROFITINFO2002",
                "EG03BJ43",
            ),
            "operating_cash": (
                "PBCEC_EG06AB_CASHINFO2007",
                "EG06BJ10",
                "PBCEC_EG05AB_CASHINFO2002",
                "EG05BJ10",
            ),
        }
        for row in rows:
            report = row["report_id"]
            personal = row["subject_type"] == "个人"
            if row["financial_conflicts"]:
                conflicts += 1
                for key in row["financial_conflicts"].split(", "):
                    assert row[key] is None
            if personal:
                assert row["cust_ind"] is None
                assert row["current_overdue_count"] is None
                continue
            for field, (
                modern_table,
                modern_field,
                old_table,
                old_field,
            ) in financial_pairs.items():
                modern = source.one(modern_table, "EA01AI01", report).get(modern_field)
                old = source.one(old_table, "EA01AI01", report).get(old_field)
                if (
                    modern is not None
                    and old is not None
                    and Decimal(str(modern)) != Decimal(str(old))
                ):
                    assert row[field] is None and field in row["financial_conflicts"].split(", ")
                elif modern is not None:
                    close(row[field], modern)
                else:
                    # Current supplied records all have the modern version; missing is not zero.
                    assert row[field] is None
            accounts = source.rows("PBCEC_ED01A_BASICLOANACCT", "EA01AI01", report)
            repay = source.rows("PBCEC_ED01B_REPAYMENT", "EA01AI01", report)
            assert len({r["ED01AI01"] for r in accounts}) == len(accounts)
            assert {r["ED01AI01"] for r in accounts} == {r["ED01AI01"] for r in repay}
            assert all(r["ED01AD07"] == "CNY" for r in accounts)
            assert len({r["ED01BR01"] for r in repay}) <= 1
            if accounts:
                multi += len(accounts) > 1
                expected = sum(Decimal(str(r["ED01BJ01"])) for r in repay)
                close(row["loan_balance_cny"], expected)
                close(row["reported_loan_balance"], expected)
                assert row["loan_account_count"] == len(accounts)
                assert row["open_loan_count"] + row["closed_loan_count"] == len(accounts)
            else:
                assert row["loan_account_count"] is None and row["loan_balance_cny"] is None
            guarantees = source.rows("PBCEC_ED04AB_GUARANTEEDETAIL", "EA01AI01", report)
            if any(r["ED04AD04"] is None for r in guarantees):
                unknown += 1
                assert row["guarantee_balance_cny"] is None
        result.update(
            financial_conflict_reports=conflicts,
            unknown_guarantee_currency_reports=unknown,
            multi_loan_account_reports=multi,
        )
    else:
        assert len({(r["unify_credit_code"], r["dt"]) for r in rows}) == len(rows)
        assert result["linked_customers"] == 1000
        by_company = {(r["unify_credit_code"], r["dt"]): r for r in rows}
        assert len(rows) == len(source.rows("T_SAIC_BASIC")) == 1002
        for base in source.rows("T_SAIC_BASIC"):
            row = by_company[base["unify_credit_code"], base["dt"]]
            # Existing capital is already in yuan. Verify no second conversion.
            for key in ("cert_capt_amt", "org_capt_amt"):
                if base[key] is not None:
                    close(row[key], base[key])
            shares = source.rows(
                "T_SAIC_SHAREHOLDER", "unify_credit_code", base["unify_credit_code"]
            )
            assert len(shares) == row["shareholder_count"]
            assert all(s["REGCAPCUR"] == "CNY" for s in shares)
            expected = sum(Decimal(str(s["SUBCONAM"])) for s in shares) * 10000
            close(row["share_subscribed_cny"], expected)
            # All mock shareholders together must match the repeated company-level total.
            for share in shares:
                close(expected, Decimal(str(share["SUMCONAM"])) * 10000)
            close(row["largest_share_ratio"], max(Decimal(str(s["FUNDEDRATIO"])) for s in shares))
            if row["exception_count"] is None:
                assert row["active_exception_count"] is None
        result["multi_shareholder_companies"] = sum(r["shareholder_count"] > 1 for r in rows)
    workbook.close()
    write_json(folder / "verified.json", result)
    print(json.dumps(dict(version=version, label=label, **result), ensure_ascii=False), flush=True)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("version", choices=BASES)
    parser.add_argument("label", choices=["征信", "工商"])
    args = parser.parse_args()
    verify(args.version, args.label)
