"""Read-only source access and conservative aggregation for import summaries."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal

import openpyxl


class Sources:
    def __init__(self, path):
        self.tables = {}
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                if sheet.title.endswith("说明"):
                    continue
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows)
                self.tables[sheet.title] = [
                    dict(
                        zip(
                            headers,
                            [
                                v.isoformat(sep=" ")
                                if isinstance(v, datetime)
                                else v.isoformat()
                                if isinstance(v, date)
                                else v
                                for v in row
                            ],
                            strict=True,
                        )
                    )
                    for row in rows
                ]
        finally:
            workbook.close()
        self.indexes = {}

    def rows(self, table, key=None, value=None):
        if key is None:
            return self.tables[table]
        if (table, key) not in self.indexes:
            index = defaultdict(list)
            for row in self.tables[table]:
                index[row[key]].append(row)
            self.indexes[table, key] = index
        return self.indexes[table, key].get(value, [])

    def one(self, table, key, value):
        rows = self.rows(table, key, value)
        if len(rows) > 1:
            raise ValueError(f"Ambiguous singleton: {table}/{key}/{value}")
        return rows[0] if rows else {}

    def lineage(self, key, value):
        refs = []
        for table in self.tables:
            if self.tables[table] and key in self.tables[table][0]:
                group = self.rows(table, key, value)
                if group:
                    refs.append(f"{table}:" + ",".join(str(int(r["source_row"])) for r in group))
        return "; ".join(refs)


def decimal(value):
    return None if value is None or value == "" else Decimal(str(value))


def total(rows, field, factor=1):
    """A missing amount invalidates the aggregate; no records means unknown."""
    values = [decimal(r.get(field)) for r in rows]
    if not values or any(v is None for v in values):
        return None
    return float(sum(values) * Decimal(factor))


def cny_total(rows, amount, currency, factor=1):
    # Unknown currencies invalidate the subtotal, rather than silently becoming CNY.
    if not rows or any(r.get(currency) in (None, "") for r in rows):
        return None
    eligible = [r for r in rows if r[currency] == "CNY"]
    return total(eligible, amount, factor) if eligible else 0


def latest_accounts(rows, account, date):
    grouped = defaultdict(list)
    for row in rows:
        if not row.get(account):
            raise ValueError(f"Missing account key {account}")
        grouped[row[account]].append(row)
    result = []
    for group in grouped.values():
        dates = [r.get(date) for r in group]
        if len(group) > 1 and any(d is None for d in dates):
            raise ValueError(f"Cannot order account snapshots: {account}/{date}")
        newest = max((str(d) for d in dates if d is not None), default=None)
        matches = [r for r in group if str(r.get(date)) == newest] if newest else group
        if len(matches) != 1:
            raise ValueError(f"Ambiguous account snapshot: {account}/{date}")
        result.append(matches[0])
    return result


def known_count(rows, predicate):
    decisions = [predicate(row) for row in rows]
    return None if not decisions or None in decisions else sum(decisions)


def count(rows):
    return len(rows) if rows else None


class Summary:
    def __init__(self, label, grain, detail, source):
        self.label, self.grain, self.detail, self.source = label, grain, detail, source
        self.fields, self.records = [], []

    def field(self, key, name, source, rule="按键读取原值；空白保留", kind="text", unit=""):
        self.fields.append(dict(key=key, name=name, source=source, rule=rule, kind=kind, unit=unit))

    def payload(self, version, source_hash):
        keys = [f["key"] for f in self.fields]
        if len(keys) != len(set(keys)):
            raise ValueError("Duplicate summary fields")
        for row in self.records:
            if set(row) - set(keys):
                raise ValueError(f"Undocumented fields: {set(row) - set(keys)}")
        return dict(
            label=self.label,
            version=version,
            grain=self.grain,
            detail=self.detail,
            source_sha256=source_hash,
            fields=self.fields,
            rows=[
                [
                    None
                    if row.get(f["key"]) is None
                    else (
                        str(row[f["key"]]) if f["kind"] == "text" else float(decimal(row[f["key"]]))
                    )
                    for f in self.fields
                ]
                for row in self.records
            ],
            source_tables=len(self.source.tables),
            source_rows=sum(map(len, self.source.tables.values())),
        )


def customer_lookup(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        records = workbook["数据"].iter_rows(values_only=True)
        header = next(records)
        result = {}
        for row in records:
            values = dict(zip(header, row, strict=True))
            code = values["unify_credit_code"]
            if not code or code in result:
                raise ValueError("Customer credit code missing or duplicated")
            result[code] = values["cust_ind"]
        return result
    finally:
        workbook.close()
