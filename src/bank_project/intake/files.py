"""Bounded CSV/XLSX readers; no database, ontology or model dependencies."""

import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException

from bank_project.intake.models import (
    Column,
    DataRow,
    IntakeError,
    Limits,
    ParsedSource,
    ParsedTable,
)
from bank_project.intake.values import cell_text

_INTEGER = re.compile(r"^[+-]?(?:0|[1-9][0-9]*)$")
_DECIMAL = re.compile(r"^[+-]?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?$")


def observed_type(value: str | None) -> str | None:
    if value in (None, ""):
        return None
    if value.lower() in {"true", "false"}:
        return "boolean"
    if _INTEGER.fullmatch(value):
        return "integer"
    if _DECIMAL.fullmatch(value):
        return "decimal"
    try:
        if len(value) == 10:
            date.fromisoformat(value)
            return "date"
        if "T" in value or " " in value:
            datetime.fromisoformat(value)
            return "datetime"
    except ValueError:
        pass
    return "text"


class TableReader:
    """Validate row shape without changing source values or inventing keys."""

    def __init__(self, name: str, headers: list[object], limits: Limits, emit=None):
        self.name = name
        self.limits = limits
        if not headers or len(headers) > limits.max_columns:
            raise IntakeError(f"表头必须包含 1–{limits.max_columns} 列", location=name)
        names = [str(value).strip() if value is not None else "" for value in headers]
        if any(not name or len(name) > 255 for name in names):
            raise IntakeError("列名不能为空或超过 255 字符，请检查表头行", location=name)
        if len(set(names)) != len(names):
            raise IntakeError("表头存在重复列名，请为每一列设置独立名称", location=name)
        self.columns = [Column(name=name) for name in names]
        self.rows: list[DataRow] = []
        self.emit, self.count = emit, 0
        self.types: list[set[str]] = [set() for _ in names]
        self.nulls = [False for _ in names]

    def add(self, number: int, values: list[object]) -> None:
        location = f"{self.name} · 第 {number} 行"
        if len(values) != len(self.columns):
            raise IntakeError("数据列数与表头不一致，请检查分隔符或表头行", location=location)
        if self.count >= self.limits.max_rows:
            raise IntakeError(
                f"单表超过 {self.limits.max_rows:,} 行上限，未保存此批次", location=location
            )
        texts = [cell_text(value, self.limits, location) for value in values]
        for i, value in enumerate(texts):
            kind = observed_type(value)
            if kind:
                self.types[i].add(kind)
            else:
                self.nulls[i] = True
        row = DataRow(number=number, values=texts)
        if len(row.model_dump_json().encode()) > 2 * 1024 * 1024:
            raise IntakeError("单行内容超过 2 MB，请拆分超长字段", location=location)
        self.count += 1
        if self.emit:
            self.emit(row)
        else:
            self.rows.append(row)

    def finish(self) -> ParsedTable:
        for i, column in enumerate(self.columns):
            kinds = self.types[i]
            if len(kinds) == 1:
                column.data_type = next(iter(kinds))
            elif kinds and kinds <= {"integer", "decimal"}:
                column.data_type = "decimal"
            elif len(kinds) > 1:
                column.data_type = "mixed"
            column.nullable = self.nulls[i] or not self.count
        return ParsedTable(self.name, self.columns, self.rows)


def _decode_csv(content: bytes, encoding: str) -> tuple[str, str]:
    candidates = ("utf-8-sig", "gb18030") if encoding == "auto" else (encoding,)
    for candidate in candidates:
        try:
            text = content.decode(candidate)
            if "\x00" in text:
                raise IntakeError("文件包含空字节，请将 CSV 另存为 UTF-8")
            return text, candidate
        except UnicodeDecodeError:
            continue
    raise IntakeError("无法按所选编码读取 CSV，请选择正确编码或另存为 UTF-8")


def _csv(
    content: bytes, filename: str, header_row: int, encoding: str, delimiter: str, limits: Limits
) -> ParsedSource:
    text, actual_encoding = _decode_csv(content, encoding)
    separators = {"comma": ",", "tab": "\t", "semicolon": ";"}
    if delimiter == "auto":
        try:
            separator = csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
        except csv.Error:
            separator = ","
    else:
        separator = separators[delimiter]
    # newline='' preserves embedded CR/LF inside quoted values.
    iterator = csv.reader(StringIO(text, newline=""), delimiter=separator, strict=True)
    reader = None
    cells = 0
    try:
        for number, values in enumerate(iterator, start=1):
            if number < header_row:
                continue
            if reader is None:
                reader = TableReader(Path(filename).stem, values, limits)
            elif values:  # ignore physically empty lines, retain rows of empty fields
                cells += len(values)
                if cells > limits.max_cells:
                    raise IntakeError("文件超过单批次单元格上限，未保存此批次")
                reader.add(number, values)
    except csv.Error as exc:
        raise IntakeError(
            "CSV 格式错误或单元格过长，请检查引号、分隔符及换行",
            location=f"记录 {iterator.line_num}",
        ) from exc
    if reader is None:
        raise IntakeError("找不到指定的表头行")
    return ParsedSource(
        [reader.finish()],
        [f"CSV 使用 {actual_encoding} 解码；所有原始文本保留，显示类型仅供参考。"],
    )


def _check_archive(archive: ZipFile, header_row: int, limits: Limits) -> None:
    entries = archive.infolist()
    if len(entries) > 10_000 or sum(item.file_size for item in entries) > limits.max_expanded_bytes:
        raise IntakeError("工作簿解压后过大，请拆分文件")
    cells = 0
    for entry in entries:
        if not re.fullmatch(r"xl/worksheets/sheet[^/]+\.xml", entry.filename):
            continue
        with archive.open(entry) as stream:
            for _, element in iterparse(stream, events=("end",)):
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "row" and int(element.get("r", "0")) > limits.max_rows + header_row:
                    raise IntakeError("Sheet 行号超出读取上限，请移除远处的空行或拆分文件")
                if tag == "c":
                    cells += 1
                    coordinate = re.fullmatch(r"([A-Z]+)([0-9]+)", element.get("r", ""))
                    if coordinate and column_index_from_string(coordinate[1]) > limits.max_columns:
                        raise IntakeError(f"Sheet 超过 {limits.max_columns} 列上限")
                    if cells > limits.max_cells:
                        raise IntakeError("工作簿超过单批次单元格上限，未保存此批次")
                element.clear()


def _xlsx(
    content: bytes, header_row: int, limits: Limits, skip_description_sheets: bool
) -> ParsedSource:
    try:
        with ZipFile(BytesIO(content)) as archive:
            _check_archive(archive, header_row, limits)
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
        raise IntakeError("无法读取 XLSX，请确认文件未损坏且没有密码保护") from exc
    tables, warnings = [], []
    cell_count = 0
    try:
        data_sheet_count = sum(
            not (
                skip_description_sheets
                and (sheet.title.endswith("说明") or sheet.title.casefold() == "readme")
            )
            for sheet in workbook.worksheets
        )
        if data_sheet_count > limits.max_tables:
            raise IntakeError(f"工作簿超过 {limits.max_tables} 个 Sheet 上限")
        for sheet in workbook.worksheets:
            if skip_description_sheets and (
                sheet.title.endswith("说明") or sheet.title.casefold() == "readme"
            ):
                warnings.append(
                    f"说明页「{sheet.title}」已跳过；如需作为数据读取，可关闭“跳过说明页”。"
                )
                continue
            reader = None
            # Do not trust cached worksheet dimensions; iterate actual XML cells.
            sheet.reset_dimensions()
            for number, row in enumerate(sheet.iter_rows(), start=1):
                cell_count += len(row)
                if cell_count > limits.max_cells:
                    raise IntakeError("工作簿超过单批次单元格上限，未保存此批次")
                if number < header_row:
                    continue
                if len(row) > limits.max_columns:
                    raise IntakeError(
                        f"Sheet 超过 {limits.max_columns} 列上限", location=sheet.title
                    )
                values = []
                for cell in row:
                    if cell.data_type in {"f", "e"}:
                        raise IntakeError(
                            "发现公式或错误单元格，请先粘贴为值后导入",
                            location=f"{sheet.title} · {cell.coordinate}",
                        )
                    values.append(cell.value)
                if reader is None:
                    while values and values[-1] is None:
                        values.pop()
                    if not values:
                        continue
                    if number != header_row:
                        raise IntakeError("指定表头行为空，请调整表头行号", location=sheet.title)
                    reader = TableReader(sheet.title, values, limits)
                elif any(value is not None for value in values):
                    width = len(reader.columns)
                    if any(value is not None for value in values[width:]):
                        raise IntakeError(
                            "数据超出表头列范围", location=f"{sheet.title} · 第 {number} 行"
                        )
                    values = values[:width] + [None] * max(0, width - len(values))
                    reader.add(number, values)
            if reader:
                tables.append(reader.finish())
            else:
                warnings.append(f"空白 Sheet「{sheet.title}」已跳过。")
    finally:
        workbook.close()
    if not tables:
        raise IntakeError("工作簿没有可导入的表")
    warnings.append("Excel 按单元格底层值读取；编号请在原文件中存为文本，显示格式不作为数据。")
    return ParsedSource(tables, warnings)


def parse_file(
    content: bytes,
    filename: str,
    limits: Limits,
    *,
    header_row: int = 1,
    encoding: str = "auto",
    delimiter: str = "auto",
    skip_description_sheets: bool = True,
) -> ParsedSource:
    if not content:
        raise IntakeError("文件为空")
    if len(content) > limits.max_upload_bytes:
        raise IntakeError("文件超过上传大小上限", status=413)
    if Path(filename).suffix.lower() == ".csv":
        return _csv(content, filename, header_row, encoding, delimiter, limits)
    if Path(filename).suffix.lower() == ".xlsx":
        try:
            return _xlsx(content, header_row, limits, skip_description_sheets)
        except (ParseError, DefusedXmlException, BadZipFile, EOFError, RuntimeError) as exc:
            raise IntakeError("工作簿结构不合法或包含不支持的内容，请重新导出 XLSX") from exc
    raise IntakeError("目前支持 .xlsx 和 .csv 文件")
