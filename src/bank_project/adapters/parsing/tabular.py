import csv
import io
from datetime import date, datetime, time, timedelta
from pathlib import PurePath
from zipfile import ZipFile

from bank_project.adapters.parsing.excel import worksheet_numbers
from bank_project.adapters.values import json_scalar
from bank_project.contracts.errors import IntakeError, LimitExceeded
from bank_project.contracts.intake import ParsedSource, SourceRow
from bank_project.contracts.schema import IntakeLimits
from bank_project.contracts.serialization import json_records, strict_json


def decode_text(content: bytes) -> str:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IntakeError("文本必须为 UTF-8；请按 UTF-8 重新导出文件") from exc
    if any(ord(c) < 32 and c not in "\r\n\t\f" for c in text):
        raise IntakeError("文本含二进制控制字符")
    return text


class TableParser:
    """Preserve any table's columns and records; business schemas are not import concerns."""

    def __init__(self, limits: IntakeLimits):
        self.limits = limits

    def headers(self, headers: list) -> None:
        if not headers:
            raise IntakeError("表头为空")
        if len(headers) > self.limits.columns:
            raise LimitExceeded("表格列数超限")
        if any(not isinstance(h, str) or not h.strip() for h in headers):
            raise IntakeError("表头必须为非空文本")
        if len(set(headers)) != len(headers):
            raise IntakeError("表头包含重复列")

    def csv(self, content: bytes, filename: str, table: str | None) -> ParsedSource:
        reader = csv.reader(io.StringIO(decode_text(content), newline=""), strict=True)
        try:
            headers = next(reader, [])
            self.headers(headers)
            rows = []
            for index, cells in enumerate(reader, 2):
                if len(cells) != len(headers):
                    raise IntakeError(f"CSV 记录 {index} 的列数与表头不一致")
                rows.append(
                    SourceRow(
                        locator=f"csv:record:{index}", values=dict(zip(headers, cells, strict=True))
                    )
                )
                self.check_rows(len(rows))
        except csv.Error as exc:
            raise IntakeError("CSV 格式无效或单元格过大") from exc
        return self.result(filename, table, rows)

    def json(self, content: bytes, filename: str, table: str | None, lines: bool) -> ParsedSource:
        text = decode_text(content)
        if lines:
            data = (
                (f"jsonl:line:{i}", line)
                for i, line in enumerate(io.StringIO(text), 1)
                if line.strip()
            )
        else:
            data = (
                (f"json:record:{i}", record)
                for i, record in enumerate(json_records(text, self.limits.rows), 1)
            )
        rows = []
        for locator, record in data:
            self.check_rows(len(rows) + 1)
            if lines:
                record = strict_json(record)
            if not isinstance(record, dict):
                raise IntakeError("每条记录必须是 JSON 对象")
            self.headers(list(record))
            rows.append(SourceRow(locator=locator, values=record))
        return self.result(filename, table, rows)

    def xlsx(
        self, content: bytes, filename: str, table: str | None, sheet: str | None
    ) -> ParsedSource:
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            if not workbook.worksheets:
                raise IntakeError("Excel 没有工作表")
            if sheet is not None and sheet not in workbook.sheetnames:
                raise IntakeError("指定的 Excel 工作表不存在")
            ws = workbook[sheet] if sheet is not None else workbook.worksheets[0]
            # openpyxl normally casts numeric XML values to binary floats. Preserve the
            # original decimal lexeme instead; Excel itself may already have rounded it.
            with ZipFile(io.BytesIO(content)) as archive:
                with archive.open(ws._worksheet_path.lstrip("/")) as xml:
                    numbers = worksheet_numbers(xml, self.limits)
            ws.reset_dimensions()
            header_cells = next(ws.iter_rows(min_row=1, max_row=1), ())
            headers = [cell.value for cell in header_cells]
            while headers and headers[-1] is None:
                headers.pop()
            self.headers(headers)
            if any(cell.data_type in {"f", "e"} for cell in header_cells):
                raise IntakeError("Excel 表头不能包含公式或错误值")
            rows = []
            for index, cells in enumerate(ws.iter_rows(min_row=2), 2):
                if index > self.limits.rows + 1:
                    raise LimitExceeded("Excel 行数超限，请拆分批次")
                if len(cells) > self.limits.columns:
                    raise LimitExceeded("Excel 列数超限")
                if any(cell.value is not None for cell in cells[len(headers) :]):
                    raise IntakeError(f"Excel 第 {index} 行存在无表头或超出列数限制的数据")
                if not any(cell.value is not None for cell in cells):
                    continue
                values = {}
                for header, cell in zip(headers, cells, strict=False):
                    if cell.data_type in {"f", "e"}:
                        raise IntakeError(f"Excel 第 {index} 行包含公式或错误值，请先导出为值")
                    value = cell.value
                    if isinstance(value, (datetime, date, time, timedelta)):
                        value = json_scalar(value)
                    elif value is not None and cell.data_type == "n":
                        value = numbers.get((cell.row, cell.column), value)
                    values[header] = value
                values = {header: values.get(header) for header in headers}
                rows.append(SourceRow(locator=f"xlsx:{ws.title}:row:{index}", values=values))
            result = self.result(filename, table or ws.title, rows)
            if sheet is None and len(workbook.sheetnames) > 1:
                result.warnings.append(
                    "默认只导入第一个工作表；其余页面未导入，导入其他页面请指定 sheet 和新的 batch_id"
                )
            return result
        finally:
            workbook.close()

    def result(self, filename: str, table: str | None, rows: list[SourceRow]) -> ParsedSource:
        if not rows:
            raise IntakeError("表中没有数据记录")
        return ParsedSource(
            filename=filename, kind="table", table=table or PurePath(filename).stem, rows=rows
        )

    def check_rows(self, count: int) -> None:
        if count > self.limits.rows:
            raise LimitExceeded(f"记录数超过 {self.limits.rows}，请拆分批次")
