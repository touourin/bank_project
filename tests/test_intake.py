"""Intake invariants: value fidelity, all-or-nothing batches and API isolation."""

import sqlite3
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from bank_project.intake.files import parse_file
from bank_project.intake.models import IntakeError, Limits
from bank_project.intake.store import BatchStore
from bank_project.main import create_app
from bank_project.settings import Settings


def workbook(sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(Settings(_env_file=None, data_dir=tmp_path))) as client:
        yield client


def upload(client, text="编号,金额\n00123,12345678901234567890.1200\n", filename="客户.csv"):
    return client.post(
        "/api/v1/intake/uploads", params={"filename": filename}, content=text.encode()
    )


def test_csv_keeps_identifiers_decimal_precision_whitespace_and_quoted_newlines():
    content = 'id,amount,note\r\n00123,12345678901234567890.1200," 中 文\r\n第二行 "\r\n'.encode(
        "utf-8-sig"
    )
    table = parse_file(content, "sample.csv", Limits()).tables[0]
    assert table.rows[0].values == ["00123", "12345678901234567890.1200", " 中 文\r\n第二行 "]
    assert table.columns[0].data_type == "text"
    assert table.columns[1].data_type == "decimal"
    assert not any(column.primary_key for column in table.columns)


@pytest.mark.parametrize("encoding", ["utf-8-sig", "gb18030"])
def test_csv_encoding_semicolon_and_header_offset(encoding):
    data = "客户数据\n客户编号;客户姓名\n00001;客户甲\n".encode(encoding)
    result = parse_file(data, "客户.csv", Limits(), header_row=2, delimiter="semicolon")
    assert result.tables[0].rows[0].values == ["00001", "客户甲"]
    assert result.tables[0].rows[0].number == 3


@pytest.mark.parametrize(
    "content",
    [
        b"id,id\n1,2\n",
        b"id,\n1,2\n",
        b"id,name\n1,2,3\n",
        b"id,name\n1\n",
        b'id,name\n1,"unterminated',
        b"",
    ],
)
def test_invalid_csv_never_silently_discards_columns(content):
    with pytest.raises(IntakeError):
        parse_file(content, "bad.csv", Limits())


def test_xlsx_multiple_sheets_and_empty_cells_keep_source_rows():
    raw = workbook(
        {
            "客户": [["id", "姓名", "备注"], ["0001", "甲", None], [], ["0002", "乙", "中文"]],
            "旅程": [["客户编号", "事件"], ["0001", "开户"]],
            "空白": [],
        }
    )
    parsed = parse_file(raw, "数据.xlsx", Limits())
    assert [t.name for t in parsed.tables] == ["客户", "旅程"]
    assert parsed.tables[0].rows[0].values == ["0001", "甲", None]
    assert parsed.tables[0].rows[1].number == 4
    assert any("空白" in warning for warning in parsed.warnings)


@pytest.mark.parametrize("rows", [[["id", "amount"], ["A", "=1+2"]], [["id"], ["A", "unheaded"]]])
def test_xlsx_rejects_formulas_and_unheaded_values(rows):
    with pytest.raises(IntakeError):
        parse_file(workbook({"客户": rows}), "bad.xlsx", Limits())


def test_xlsx_header_selection():
    parsed = parse_file(
        workbook({"客户": [["说明"], ["id", "name"], ["001", "甲"]]}),
        "a.xlsx",
        Limits(),
        header_row=2,
    )
    assert parsed.tables[0].rows[0].number == 3


def test_parser_limits_are_errors_not_silent_truncation():
    with pytest.raises(IntakeError, match="行上限"):
        parse_file(b"id\n1\n2\n", "a.csv", Limits(max_rows=1))
    with pytest.raises(IntakeError, match="单元格"):
        parse_file(b"a,b\n1,2\n", "a.csv", Limits(max_cells=1))
    with pytest.raises(IntakeError, match="解压"):
        parse_file(workbook({"客户": [["id"], ["001"]]}), "a.xlsx", Limits(max_expanded_bytes=10))
    with pytest.raises(IntakeError, match="读取 XLSX"):
        parse_file(b"broken", "a.xlsx", Limits())


def test_batch_duplicate_names_pagination_and_restart(tmp_path):
    store = BatchStore(tmp_path / "staging.sqlite3")
    data = parse_file(b"id\n001\n002\n", "customers.csv", Limits())
    first = store.save(data, name="a.csv", source_kind="file", source="a.csv")
    second = store.save(data, name="a.csv", source_kind="file", source="a.csv")
    assert first.id != second.id
    restored = BatchStore(store.path)
    assert restored.list(0, 1).total == 2
    assert len(restored.list(0, 1).items) == 1
    preview = restored.preview(first.id, first.tables[0].id, 1, 1)
    assert preview.rows[0].values == ["002"]
    with pytest.raises(IntakeError):
        restored.preview(second.id, first.tables[0].id, 0, 10)
    restored.delete(first.id)
    assert restored.get(second.id).row_count == 2
    with restored.connect() as db:
        assert db.execute("SELECT count(*) FROM rows").fetchone()[0] == 2


def test_failed_database_write_rolls_back_whole_batch(tmp_path):
    store = BatchStore(tmp_path / "db.sqlite3")
    with store.connect() as db:
        db.execute("""CREATE TRIGGER fail_second BEFORE INSERT ON tables
            WHEN json_extract(new.metadata, '$.name') = 'bad'
            BEGIN SELECT RAISE(ABORT, 'simulated disk failure'); END""")
    parsed = parse_file(
        workbook({"good": [["id"], ["1"]], "bad": [["id"], ["2"]]}), "a.xlsx", Limits()
    )
    with pytest.raises(sqlite3.IntegrityError):
        store.save(parsed, name="a", source_kind="file", source="a")
    assert store.list(0, 10).total == 0
    with store.connect() as db:
        assert db.execute("SELECT count(*) FROM rows").fetchone()[0] == 0


def test_concurrent_uploads_do_not_replace_each_other(tmp_path):
    store = BatchStore(tmp_path / "db.sqlite3")
    parsed = parse_file(b"id\n1", "a.csv", Limits())

    def save(_):
        return store.save(parsed, name="same", source_kind="file", source="same").id

    with ThreadPoolExecutor(max_workers=4) as pool:
        ids = list(pool.map(save, range(8)))
    assert len(set(ids)) == 8
    assert store.list(0, 10).total == 8


def test_api_upload_preview_failure_and_delete(client):
    response = upload(client)
    assert response.status_code == 201
    batch = response.json()
    table_id = batch["tables"][0]["id"]
    path = f"/api/v1/intake/batches/{batch['id']}"
    assert client.get(path).json()["tables"][0]["name"] == "客户"
    preview = client.get(f"{path}/tables/{table_id}").json()
    assert preview["rows"][0]["values"] == ["00123", "12345678901234567890.1200"]
    assert upload(client, "id,id\n1,2").status_code == 422
    assert client.get("/api/v1/intake/batches").json()["total"] == 1
    assert client.delete(path).status_code == 204
    assert client.get(path).status_code == 404


def test_auth_origin_and_invalid_password_payload_are_safe(tmp_path):
    settings = Settings(
        _env_file=None, data_dir=tmp_path, api_token="a-very-long-private-api-token"
    )
    with TestClient(create_app(settings)) as client:
        assert client.get("/health").status_code == 200
        assert upload(client).status_code == 401
        client.headers["Authorization"] = "Bearer a-very-long-private-api-token"
        assert upload(client).status_code == 201
        assert (
            client.get(
                "/api/v1/intake/batches", headers={"Origin": "https://untrusted.example"}
            ).status_code
            == 403
        )
        response = client.post(
            "/api/v1/intake/mysql/catalog",
            json={
                "host": "x",
                "port": -1,
                "password": "never-echo-me",
                "surprise": "never-echo-me",
            },
        )
        assert response.status_code == 422
        assert "never-echo-me" not in response.text


def test_request_limit_counts_stream_even_without_content_length(tmp_path):
    with TestClient(
        create_app(Settings(_env_file=None, data_dir=tmp_path, intake_max_upload_mb=1))
    ) as client:
        response = client.post(
            "/api/v1/intake/uploads?filename=big.csv", content=(b"a" * 65536 for _ in range(17))
        )
        assert response.status_code == 413
        assert client.get("/api/v1/intake/batches").json()["total"] == 0
        assert client.post("/api/v1/intake/mysql/catalog", content=b"a" * 65537).status_code == 413


def test_description_sheets_are_skipped_explicitly_and_can_be_included():
    raw = workbook(
        {"数据": [["id"], ["001"]], "字段说明": [["field", "description"], ["id", "编号"]]}
    )
    parsed = parse_file(raw, "a.xlsx", Limits())
    assert len(parsed.tables) == 1
    assert any("字段说明" in warning for warning in parsed.warnings)
    parsed = parse_file(raw, "a.xlsx", Limits(), skip_description_sheets=False)
    assert len(parsed.tables) == 2


def test_configured_mysql_is_resolved_server_side_and_never_returned(tmp_path, monkeypatch):
    from bank_project.intake.mysql import MysqlSource

    captured = []

    def catalog(self, connection):
        captured.append(connection)
        return []

    monkeypatch.setattr(MysqlSource, "catalog", catalog)
    settings = Settings(_env_file=None, data_dir=tmp_path, mysql_password="configured-secret")
    with TestClient(create_app(settings)) as client:
        response = client.post("/api/v1/intake/mysql/catalog", json={})
        assert response.status_code == 200
        assert "configured-secret" not in response.text
        assert captured[0].password.get_secret_value() == "configured-secret"


def test_sparse_or_malicious_xlsx_is_rejected_before_row_expansion():
    raw = workbook({"数据": [["id"], ["001"]]})
    source, target = BytesIO(raw), BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w") as modified:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'<row r="2">', b'<row r="999999999">')
            modified.writestr(entry, content)
    with pytest.raises(IntakeError, match="行号"):
        parse_file(target.getvalue(), "sparse.xlsx", Limits())
