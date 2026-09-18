import io
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from bank_project.adapters.mappings import load_mappings
from bank_project.adapters.model_client.compatible import DisabledModel
from bank_project.adapters.parsing.parser import FileParser
from bank_project.adapters.raw_store.filesystem import FileStore
from bank_project.adapters.sources.files import InboxReader
from bank_project.adapters.sources.router import SourceRouter
from bank_project.contracts.document import DocumentExtraction
from bank_project.contracts.errors import (
    Conflict,
    DependencyUnavailable,
    IntakeError,
    LimitExceeded,
    NotFound,
)
from bank_project.contracts.intake import ExtractionRequest, RawInput
from bank_project.contracts.mapping import MappingCatalog, MappingProfile
from bank_project.contracts.models import ImportRequest
from bank_project.contracts.schema import IntakeLimits
from bank_project.extraction.service import ExtractionService
from bank_project.ingestion.service import IngestionService

ROOT = Path(__file__).resolve().parents[1]


def services(tmp_path, model=None, catalog=None):
    store = FileStore(tmp_path / "store")
    limits = IntakeLimits()
    ingestion = IngestionService(
        SourceRouter(InboxReader(tmp_path, limits.file_bytes)), FileParser(limits), store, store
    )
    extraction = ExtractionService(
        store,
        store,
        catalog or load_mappings(ROOT / "configs/mappings"),
        model or DisabledModel(),
        store,
    )
    return ingestion, extraction, store


def upload(ingestion, content, filename="other.csv", batch="one", dataset="data", **kwargs):
    if isinstance(content, str):
        content = content.encode("utf-8")
    request = ImportRequest(
        dataset_id=dataset,
        batch_id=batch,
        source_system="source",
        source_uri=f"upload:{filename}",
        **kwargs,
    )
    return ingestion.ingest(request, RawInput(filename=filename, content=content))


def order_mapping():
    return MappingProfile.model_validate(
        {
            "id": "sales_orders",
            "table": "orders",
            "match_columns": ["order_id", "buyer", "seller", "amount", "day"],
            "primary_key": ["order_id"],
            "columns": [
                {"name": "amount", "sql_type": "decimal(26,8)"},
                {"name": "day", "sql_type": "varchar(8)", "format": "YYYYMMDD"},
            ],
            "entities": [
                {
                    "alias": "buyer",
                    "entity_type": "company",
                    "keys": {"company_id": "buyer"},
                    "identity_scope": "key",
                },
                {
                    "alias": "seller",
                    "entity_type": "company",
                    "keys": {"company_id": "seller"},
                    "identity_scope": "key",
                },
            ],
            "event": {
                "event_type": "purchase",
                "date_field": "day",
                "participants": [
                    {"entity": "buyer", "role": "buyer"},
                    {"entity": "seller", "role": "seller"},
                ],
            },
            "relations": [
                {
                    "subject": "buyer",
                    "predicate": "buys_from",
                    "object": "seller",
                    "link_event": True,
                }
            ],
        }
    )


def test_any_table_works_without_bank_schema_or_model(tmp_path):
    ingestion, extraction, store = services(tmp_path, catalog=MappingCatalog(profiles=()))
    receipt = upload(ingestion, 'sku,备注\n00007,"中文,保留\n第二行"\n00007,另一条\n')
    assert receipt.records == 2
    summary = extraction.extract(ExtractionRequest(dataset_id="data", batch_id="one"))
    batch = extraction.result("data", "one")
    assert summary.entities == 2 and summary.events == 0
    assert batch.entities[0].properties == {"sku": "00007", "备注": "中文,保留\n第二行"}
    assert batch.entities[0].candidate_id != batch.entities[1].candidate_id
    assert store.read(receipt.artifact).decode().startswith("sku,备注")
    assert batch.evidence[0].fields["sku"] == "00007"


def test_new_business_requires_only_mapping_configuration(tmp_path):
    profiles = tmp_path / "profiles"
    profiles.mkdir()
    (profiles / "orders.json").write_text(order_mapping().model_dump_json())
    ingestion, extraction, _ = services(tmp_path, catalog=load_mappings(profiles))
    upload(
        ingestion,
        "order_id,buyer,seller,amount,day\n001,A,B,99999999999999999.12345678,20260901\n002,A,B,200.00,20260902",
        filename="orders.csv",
    )
    result = extraction.extract(ExtractionRequest(dataset_id="data", batch_id="one"))
    batch = extraction.result("data", "one")
    assert (result.entities, result.events, result.relations) == (2, 2, 2)
    assert len({e.event_id for e in batch.events}) == 2
    assert batch.events[0].properties["source_fields"]["amount"] == "99999999999999999.12345678"
    assert str(batch.events[0].occurred_on) == "2026-09-01"
    assert batch.events[0].occurred_at is None
    buyer = next(e for e in batch.entities if e.external_keys[0].value == "A")
    seller = next(e for e in batch.entities if e.external_keys[0].value == "B")
    assert all(
        r.subject_id == buyer.candidate_id and r.object_id == seller.candidate_id
        for r in batch.relations
    )


def test_batch_retries_are_idempotent_and_different_content_conflicts(tmp_path):
    ingestion, extraction, _ = services(tmp_path)
    first = upload(ingestion, "x\n1\n")
    assert upload(ingestion, "x\n1\n") == first
    with pytest.raises(Conflict):
        upload(ingestion, "x\n2\n")
    request = ExtractionRequest(dataset_id="data", batch_id="one")
    converted = extraction.extract(request)
    assert extraction.extract(request) == converted
    restored_ingestion, restored_extraction, _ = services(tmp_path)
    assert restored_ingestion.get("data", "one") == first
    assert restored_extraction.state("data", "one").status == "completed"
    assert restored_extraction.extract(request) == converted
    with pytest.raises(NotFound):
        restored_extraction.result("another-dataset", "one")


def test_duplicate_business_key_fails_without_publishing_partial_results(tmp_path):
    ingestion, extraction, store = services(
        tmp_path, catalog=MappingCatalog(profiles=(order_mapping(),))
    )
    upload(
        ingestion,
        "order_id,buyer,seller,amount,day\n1,A,B,20,20260901\n1,A,B,30,20260902",
        filename="orders.csv",
    )
    with pytest.raises(IntakeError, match="重复主键"):
        extraction.extract(ExtractionRequest(dataset_id="data", batch_id="one"))
    assert store.load_extraction("data", "one") is None
    assert extraction.state("data", "one").status == "failed"
    assert store.load_import("data", "one") is not None


def test_changed_mapping_never_silently_overwrites_completed_result(tmp_path):
    ingestion, extraction, store = services(
        tmp_path, catalog=MappingCatalog(profiles=(order_mapping(),))
    )
    upload(ingestion, "order_id,buyer,seller,amount,day\n1,A,B,20,20260901", filename="orders.csv")
    req = ExtractionRequest(dataset_id="data", batch_id="one")
    extraction.extract(req)
    changed = order_mapping().model_copy(update={"version": "2"})
    other = ExtractionService(
        store, store, MappingCatalog(profiles=(changed,)), DisabledModel(), store
    )
    with pytest.raises(Conflict, match="配置已变化"):
        other.extract(req)


def test_lock_and_corrupt_raw_source_are_detected(tmp_path):
    ingestion, extraction, store = services(tmp_path)
    with store.lock("data", "one"):
        with pytest.raises(Conflict):
            upload(ingestion, "x\n1")
    receipt = upload(ingestion, "x\n1")
    (store.root / "raw" / f"{receipt.artifact.digest}.bin").write_bytes(b"tampered")
    with pytest.raises(DependencyUnavailable, match="摘要校验失败"):
        extraction.extract(ExtractionRequest(dataset_id="data", batch_id="one"))
    assert extraction.state("data", "one").status == "failed"


@pytest.mark.parametrize(
    "filename,content",
    [
        ("test.csv", b"a,a\n1,2"),
        ("test.csv", b"a,b\n1"),
        ("test.csv", b'a\n"unterminated'),
        ("test.csv", b"a\n\xff"),
        ("test.json", b'[{"a":1,"a":2}]'),
        ("test.json", b'[{"a":NaN}]'),
        ("test.json", b'[{"a":Infinity}]'),
        ("test.docx", b"not a zip"),
        ("test.exe", b"anything"),
        ("../escape.csv", b"a\n1"),
    ],
)
def test_invalid_inputs_have_no_committed_import(tmp_path, filename, content):
    ingestion, _, store = services(tmp_path)
    with pytest.raises(IntakeError):
        upload(ingestion, content, filename=filename)
    assert store.load_import("data", "one") is None


def test_limits_and_source_path_escapes(tmp_path):
    limits = replace(IntakeLimits(), file_bytes=8, rows=1)
    request = ImportRequest(
        dataset_id="d", batch_id="b", source_system="s", source_uri="file:x.csv"
    )
    parser = FileParser(limits)
    with pytest.raises(LimitExceeded):
        parser.parse(RawInput(filename="x.csv", content=b"x\n123456789"), request)
    with pytest.raises(LimitExceeded):
        parser.parse(RawInput(filename="x.csv", content=b"x\n1\n2"), request)
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    secret = tmp_path / "secret.csv"
    secret.write_text("secret")
    (inbox / "linked.csv").symlink_to(secret)
    (inbox / "directory").symlink_to(tmp_path, target_is_directory=True)
    reader = InboxReader(inbox, 100)
    for uri in [
        "file:../secret.csv",
        "file:/etc/passwd",
        "file:linked.csv",
        "file:directory/secret.csv",
        "https://example.com/data",
    ]:
        with pytest.raises(IntakeError):
            reader.read(request.model_copy(update={"source_uri": uri}))


def test_excel_sheet_selection_formulas_and_raw_rows(tmp_path):
    book = Workbook()
    book.active.title = "inventory"
    book.active.append(["sku", "name"])
    book.active.append(["0001", "设备"])
    notes = book.create_sheet("字段说明")
    notes.append(["字段", "含义"])
    notes.append(["sku", "编号"])
    out = io.BytesIO()
    book.save(out)
    ingestion, extraction, _ = services(tmp_path)
    receipt = upload(ingestion, out.getvalue(), filename="inventory.xlsx")
    assert receipt.records == 1 and receipt.warnings
    extraction.extract(ExtractionRequest(dataset_id="data", batch_id="one"))
    assert extraction.result("data", "one").entities[0].properties["sku"] == "0001"
    book.active["A2"] = "=1+1"
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(IntakeError, match="公式"):
        upload(ingestion, out.getvalue(), filename="formula.xlsx", batch="formula")


def test_docx_paragraph_provenance_and_xml_entity_rejection(tmp_path):
    ingestion, _, store = services(tmp_path)
    out = io.BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>机器检修完成。</w:t></w:r></w:p></w:body></w:document>',
        )
    upload(ingestion, out.getvalue(), filename="report.docx")
    parsed = store.load_import("data", "one").parsed
    assert parsed.blocks[0].text == "机器检修完成。"
    assert parsed.blocks[0].locator == "docx:paragraph:1"
    out = io.BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "word/document.xml", '<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>'
        )
    with pytest.raises(IntakeError):
        upload(ingestion, out.getvalue(), filename="evil.docx", batch="evil")


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_full_bank_fixtures_are_only_two_profiles_of_the_generic_engine(tmp_path, extension):
    ingestion, extraction, _ = services(tmp_path)
    files = [
        ("CCM_C_CUST_FLAG_INFO.csv" if extension == "csv" else "客户标签.xlsx", "tags", 1000, 0),
        (
            "E_CRM_C_CUST_TOUR_EVT_SUM.csv" if extension == "csv" else "客户旅程.xlsx",
            "journey",
            2098,
            16700,
        ),
    ]
    for filename, batch, entities, events in files:
        upload(
            ingestion,
            (ROOT / "examples/mock" / filename).read_bytes(),
            filename=filename,
            batch=batch,
        )
        summary = extraction.extract(ExtractionRequest(dataset_id="data", batch_id=batch))
        assert (summary.entities, summary.events) == (entities, events)
    tags = extraction.result("data", "tags")
    journey = extraction.result("data", "journey")
    tag_customers = {e.candidate_id for e in tags.entities}
    assert all(event.participants[0].entity_id in tag_customers for event in journey.events)
    assert len({e.event_id for e in journey.events}) == 16700
    assert all(e.occurred_at is None and e.occurred_on is not None for e in journey.events)
    unknown = next(e for e in journey.events if e.event_type == "YWJC0017")
    assert unknown.properties["event_name"] == "未定义1"
    assert unknown.properties["requires_confirmation"] is True
    assert any(e.event_type == "MOCK_CREDIT_SIGN" for e in journey.events)


class FlakyModel:
    fingerprint = "flaky-test"
    calls = 0
    fail = True

    def extract(self, text):
        self.calls += 1
        if self.calls == 2 and self.fail:
            raise DependencyUnavailable("test failure")
        return DocumentExtraction(entities=[], events=[], relations=[])


def test_model_failure_can_resume_cached_chunks_without_partial_output(tmp_path):
    model = FlakyModel()
    ingestion, extraction, store = services(tmp_path, model=model)
    upload(ingestion, "文" * 8000, filename="report.txt")
    request = ExtractionRequest(dataset_id="data", batch_id="one")
    with pytest.raises(DependencyUnavailable):
        extraction.extract(request)
    assert store.load_extraction("data", "one") is None
    model.fail = False
    extraction.extract(request)
    assert model.calls == 4  # three chunks: cached first, retry second, then third
    assert extraction.state("data", "one").status == "completed"


@pytest.mark.parametrize(
    "extension,content",
    [
        ("json", b'[{"x":1},{"x":2},THIS_PART_MUST_NOT_BE_PARSED]'),
        ("jsonl", b'{"x":1}\n{"x":2}\nTHIS_PART_MUST_NOT_BE_PARSED'),
    ],
)
def test_json_record_limit_precedes_decoding_the_rest(extension, content):
    parser = FileParser(replace(IntakeLimits(), rows=1))
    request = ImportRequest(
        dataset_id="d", batch_id="b", source_system="s", source_uri="upload:test"
    )
    with pytest.raises(LimitExceeded):
        parser.parse(RawInput(filename=f"test.{extension}", content=content), request)


@pytest.mark.parametrize(
    "content", [b'[{"x":1},]', b'[{"x":1}]garbage', b'{"x":"\\ud800"}', b'[{"x":1} {"x":2}]']
)
def test_json_edge_cases_are_rejected(tmp_path, content):
    ingestion, _, _ = services(tmp_path)
    with pytest.raises(IntakeError):
        upload(ingestion, content, filename="broken.json")


def test_equivalent_csv_and_xlsx_keep_business_event_identity(tmp_path):
    ingestion, extraction, _ = services(
        tmp_path, catalog=MappingCatalog(profiles=(order_mapping(),))
    )
    upload(
        ingestion,
        "order_id,buyer,seller,amount,day\n001,A,B,20.01,20260901",
        filename="orders.csv",
        batch="csv",
    )
    workbook = Workbook()
    workbook.active.title = "数据"
    workbook.active.append(["order_id", "buyer", "seller", "amount", "day"])
    workbook.active.append(["001", "A", "B", "20.01", "20260901"])
    content = io.BytesIO()
    workbook.save(content)
    upload(ingestion, content.getvalue(), filename="different-name.xlsx", batch="xlsx")
    for batch in ["csv", "xlsx"]:
        extraction.extract(ExtractionRequest(dataset_id="data", batch_id=batch))
    csv = extraction.result("data", "csv")
    xlsx = extraction.result("data", "xlsx")
    assert csv.events[0].event_id == xlsx.events[0].event_id
    assert csv.relations[0].relation_id == xlsx.relations[0].relation_id
    assert csv.evidence[0].source_id != xlsx.evidence[0].source_id


def test_sparse_excel_rows_preserve_empty_columns_and_exact_decimal_xml(tmp_path):
    workbook = Workbook()
    workbook.active.append(["id", "optional", "amount"])
    workbook.active.append(["001", None, 2.25])
    workbook.active.append(["002"])
    content = io.BytesIO()
    workbook.save(content)
    ingestion, extraction, _ = services(tmp_path)
    upload(ingestion, content.getvalue(), filename="sparse.xlsx")
    extraction.extract(ExtractionRequest(dataset_id="data", batch_id="one"))
    rows = extraction.result("data", "one").entities
    assert rows[0].properties == {"id": "001", "optional": None, "amount": "2.25"}
    assert rows[1].properties == {"id": "002", "optional": None, "amount": None}


@pytest.mark.parametrize("iso_dates", [True, False])
def test_excel_temporal_values_preserve_clock_time_and_duration(tmp_path, iso_dates):
    workbook = Workbook(iso_dates=iso_dates)
    workbook.active.append(["id", "created_at", "day", "clock", "duration"])
    workbook.active.append(
        [
            "001",
            datetime(2026, 9, 18, 13, 45, 52, 123000),
            date(2026, 9, 18),
            time(13, 45, 52, 120000),
            timedelta(hours=27, seconds=12),
        ]
    )
    content = io.BytesIO()
    workbook.save(content)
    ingestion, extraction, store = services(tmp_path)
    upload(ingestion, content.getvalue(), filename="times.xlsx")
    expected = {
        "id": "001",
        "created_at": "2026-09-18T13:45:52.123000",
        "day": "2026-09-18" if iso_dates else "2026-09-18T00:00:00",
        "clock": "13:45:52.120000",
        "duration": "PT97212S",
    }
    assert store.load_import("data", "one").parsed.rows[0].values == expected
    extraction.extract(ExtractionRequest(dataset_id="data", batch_id="one"))
    assert extraction.result("data", "one").entities[0].properties == expected


def workbook_with_xml(sheet_xml):
    original = io.BytesIO()
    Workbook().save(original)
    output = io.BytesIO()
    with ZipFile(original) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = (
                    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                    f"<sheetData>{sheet_xml}</sheetData></worksheet>"
                ).encode()
            target.writestr(info, data)
    return output.getvalue()


@pytest.mark.parametrize(
    "sheet_xml,error",
    [
        ('<row r="1"><c r="XFD1"><v>1</v></c></row>', LimitExceeded),
        ('<row r="1000000"><c r="A1000000"><v>1</v></c></row>', LimitExceeded),
        ('<row r="1"><c r="A1"><v>1</v></c><c r="A1"><v>2</v></c></row>', IntakeError),
        ('<row r="1"><c r="A2"><v>1</v></c></row>', IntakeError),
        ('<row r="1"/><row r="1"/>', IntakeError),
    ],
)
def test_unsafe_excel_coordinates_fail_before_cell_arrays_are_allocated(
    tmp_path, monkeypatch, sheet_xml, error
):
    content = workbook_with_xml(sheet_xml)

    def materialization_forbidden(*args, **kwargs):
        pytest.fail("Unsafe coordinates reached openpyxl cell array allocation")

    monkeypatch.setattr(
        "openpyxl.worksheet._read_only.ReadOnlyWorksheet.iter_rows", materialization_forbidden
    )
    ingestion, _, store = services(tmp_path)
    with pytest.raises(error):
        upload(ingestion, content, filename="bad.xlsx")
    assert store.load_import("data", "one") is None


def test_shared_entities_keep_every_record_evidence_once(tmp_path):
    ingestion, extraction, _ = services(
        tmp_path, catalog=MappingCatalog(profiles=(order_mapping(),))
    )
    # The same company plays both roles in every row. No event or provenance can disappear.
    content = "order_id,buyer,seller,amount,day\n" + "\n".join(
        f"{i},A,A,20,20260918" for i in range(1000)
    )
    upload(ingestion, content, filename="orders.csv")
    extraction.extract(ExtractionRequest(dataset_id="data", batch_id="one"))
    result = extraction.result("data", "one")
    assert len(result.entities) == 1
    assert len(result.events) == len(result.relations) == len(result.evidence) == 1000
    assert result.entities[0].evidence_ids == [e.evidence_id for e in result.evidence]
